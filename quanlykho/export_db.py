import os
import sys
import sqlite3

sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

db_path = r'c:\QuanLyTonKhoNongSan\quanlykho\db.sqlite3'
excel_path = r'c:\QuanLyTonKhoNongSan\DuLieu_CoSoDuLieu_QuanLyKhoNongSan.xlsx'
sql_path = r'c:\QuanLyTonKhoNongSan\database_dump.sql'

conn = sqlite3.connect(db_path)
conn.row_factory = sqlite3.Row
cur = conn.cursor()

# 1. Export SQL Dump
print("Đang xuất file SQL dump...")
with open(sql_path, 'w', encoding='utf-8') as f:
    for line in conn.iterdump():
        f.write(f"{line}\n")
print(f"-> Đã xuất SQL dump thành công: {sql_path}")

# 2. Export Excel Workbook
print("Đang xuất file Excel đa trang (Multi-sheet)...")
wb = openpyxl.Workbook()
wb.remove(wb.active)  # remove default sheet

header_fill = PatternFill(start_color='1E6B3E', end_color='1E6B3E', fill_type='solid')
header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
data_font = Font(name='Segoe UI', size=10, color='1F2937')
thin_border = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB')
)

tables_mapping = [
    (
        '1_NongSan',
        '''
        SELECT p.id, p.code, p.name, c.name as category_name, p.unit, p.price,
               p.origin, p.quality_standard, s.name as supplier_name, p.season,
               p.safety_stock, p.storage_temp, p.strategy, p.description
        FROM core_product p
        LEFT JOIN core_category c ON p.category_id=c.id
        LEFT JOIN core_supplier s ON p.supplier_id=s.id
        ORDER BY p.id
        ''',
        ['ID', 'Mã Nông Sản', 'Tên Nông Sản', 'Nhóm Hàng', 'ĐVT', 'Giá Bán Lẻ (VNĐ)',
         'Vùng Trồng / Xuất Xứ', 'Tiêu Chuẩn Chất Lượng', 'Cơ Sở Nhập / Nhà Cung Cấp',
         'Mùa Vụ', 'Tồn An Toàn', 'Nhiệt Độ Bảo Quản', 'Chiến Lược Xuất Kho', 'Mô Tả Sản Phẩm']
    ),
    (
        '2_LoHang_FEFO',
        '''
        SELECT b.id, b.batch_code, p.name as product_name,
               COALESCE(l.zone || ' - ' || l.shelf || ' - ' || l.level, 'Kho chung') as location_str,
               b.mfg_date, b.exp_date, b.qty_initial, b.qty_current, b.unit_cost, b.status
        FROM core_batch b
        JOIN core_product p ON b.product_id=p.id
        LEFT JOIN core_location l ON b.location_id=l.id
        ORDER BY b.exp_date ASC
        ''',
        ['ID', 'Mã Lô Hàng', 'Tên Nông Sản', 'Vị Trí Kệ Kho', 'Ngày Sản Xuất (MFG)',
         'Hạn Sử Dụng (EXP)', 'Số Lượng Nhập Ban Đầu', 'Số Lượng Tồn Hiện Tại',
         'Giá Vốn Nhập (VNĐ)', 'Trạng Thái Lô']
    ),
    (
        '3_PhieuNhapKho',
        '''
        SELECT r.id, r.code, s.name as supplier_name, COALESCE(u.username, 'admin') as creator,
               r.created_at, r.total_amount, r.note
        FROM core_receipt r
        JOIN core_supplier s ON r.supplier_id=s.id
        LEFT JOIN auth_user u ON r.created_by_id=u.id
        ORDER BY r.created_at DESC
        ''',
        ['ID', 'Mã Phiếu Nhập', 'Nhà Cung Cấp / Hợp Tác Xã', 'Người Lập Phiếu',
         'Ngày Giờ Nhập Kho', 'Tổng Tiền Nhập (VNĐ)', 'Ghi Chú Nhập']
    ),
    (
        '4_PhieuXuatKho',
        '''
        SELECT i.id, i.code, COALESCE(p.name, 'Khách Lẻ / Đại Lý') as partner_name,
               COALESCE(u.username, 'admin') as creator, i.created_at, i.status,
               i.total_amount, i.note
        FROM core_issue i
        LEFT JOIN core_partner p ON i.partner_id=p.id
        LEFT JOIN auth_user u ON i.created_by_id=u.id
        ORDER BY i.created_at DESC
        ''',
        ['ID', 'Mã Phiếu Xuất', 'Khách Hàng / Đối Tác Nhận', 'Người Xuất Kho',
         'Thời Gian Xuất', 'Trạng Thái Xuất', 'Tổng Giá Trị (VNĐ)', 'Ghi Chú Xuất']
    ),
    (
        '5_ChiTietXuatKho',
        '''
        SELECT d.id, i.code as issue_code, b.batch_code, p.name as product_name,
               d.quantity, d.unit_price
        FROM core_issuedetail d
        JOIN core_issue i ON d.issue_id=i.id
        JOIN core_batch b ON d.batch_id=b.id
        JOIN core_product p ON b.product_id=p.id
        ORDER BY d.id
        ''',
        ['ID', 'Mã Phiếu Xuất', 'Lô Hàng Xuất', 'Nông Sản', 'Số Lượng Xuất', 'Đơn Giá Xuất']
    ),
    (
        '6_DonDatHang_Online',
        '''
        SELECT o.id, o.code, o.customer_name, o.customer_phone, o.shipping_address,
               o.payment_method, o.payment_status, o.order_status, o.total_amount,
               o.created_at, o.note
        FROM core_customerorder o
        ORDER BY o.created_at DESC
        ''',
        ['ID', 'Mã Đơn Hàng', 'Họ Tên Khách Hàng', 'Số Điện Thoại', 'Địa Chỉ Giao Hàng',
         'Phương Thức Thanh Toán', 'Trạng Thái Thanh Toán', 'Trạng Thái Đơn Hàng',
         'Tổng Tiền Đơn (VNĐ)', 'Thời Gian Đặt Hàng', 'Ghi Chú Khách Hàng']
    ),
    (
        '7_NhaCungCap_HTX',
        '''
        SELECT id, name, phone, email, address, notes FROM core_supplier ORDER BY id
        ''',
        ['ID', 'Tên Cơ Sở Nhập / HTX Nông Nghiệp', 'Số Điện Thoại', 'Email', 'Địa Chỉ Vùng Trồng', 'Ghi Chú']
    ),
    (
        '8_KhachHang_DaiLy',
        '''
        SELECT id, name, contact_person, phone, address, notes FROM core_partner ORDER BY id
        ''',
        ['ID', 'Tên Khách Hàng / Đại Lý', 'Người Liên Hệ', 'Số Điện Thoại', 'Địa Chỉ', 'Ghi Chú']
    ),
    (
        '9_ViTriKho',
        '''
        SELECT id, zone, shelf, level, note FROM core_location ORDER BY id
        ''',
        ['ID', 'Khu Vực Kho (Zone)', 'Kệ (Shelf)', 'Tầng (Level)', 'Ghi Chú']
    ),
    (
        '10_NhomNongSan',
        '''
        SELECT id, code, name, storage_type, description FROM core_category ORDER BY id
        ''',
        ['ID', 'Mã Nhóm', 'Tên Nhóm Nông Sản', 'Phân Loại Kho Lưu Trữ', 'Mô Tả']
    ),
    (
        '11_BaoPhe_HaoHut',
        '''
        SELECT w.id, w.code, b.batch_code, COALESCE(u.username, 'admin') as reporter,
               w.report_date, w.quantity, w.reason, w.note
        FROM core_wastereport w
        JOIN core_batch b ON w.batch_id=b.id
        LEFT JOIN auth_user u ON w.reported_by_id=u.id
        ORDER BY w.report_date DESC
        ''',
        ['ID', 'Mã Biên Bản', 'Mã Lô Hàng Bị Hao Hụt', 'Người Lập Biên Bản',
         'Ngày Ghi Nhận', 'Số Lượng Hao Hụt (kg)', 'Nguyên Nhân (Dập nát / Biến chất / Hết hạn)', 'Ghi Chú']
    ),
    (
        '12_TaiKhoan_PhanQuyen',
        '''
        SELECT u.id, u.username, COALESCE(p.full_name, u.username) as full_name,
               u.email, COALESCE(p.phone, '') as phone, COALESCE(p.address, '') as address,
               COALESCE(p.role, 'STAFF') as role,
               CASE WHEN u.is_active = 1 THEN 'Hoạt động' ELSE 'Đã khóa' END as status_str,
               u.date_joined, u.last_login
        FROM auth_user u
        LEFT JOIN core_userprofile p ON u.id=p.user_id
        ORDER BY u.id
        ''',
        ['ID', 'Tên Đăng Nhập (Username)', 'Họ Và Tên', 'Email', 'Số Điện Thoại',
         'Địa Chỉ', 'Vai Trò (Role)', 'Trạng Thái Tài Khoản', 'Ngày Tạo', 'Lần Đăng Nhập Cuối']
    )
]

for sheet_title, query, headers in tables_mapping:
    ws = wb.create_sheet(title=sheet_title)
    ws.append(headers)
    
    # Format Header
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # Insert Data
    cur.execute(query)
    rows = cur.fetchall()
    for row in rows:
        row_data = [str(item) if item is not None else '' for item in row]
        ws.append(row_data)

    # Format Data cells
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

    # Auto-adjust column width
    for col in ws.columns:
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

wb.save(excel_path)
print(f"-> Đã xuất Excel đa trang thành công: {excel_path}")
print("HOÀN TẤT XUẤT DATABASE THÀNH CÔNG!")
